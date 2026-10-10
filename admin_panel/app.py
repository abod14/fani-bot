# لوحة تحكم بوت «فني» — صفحة ويب بسيطة (Flask) بكلمة مرور، تقرأ/تكتب نفس قاعدة
# بيانات SQLite اللي يستخدمها البوت مباشرة، بدون أي طبقة تزامن إضافية.
#
# التشغيل: python admin_panel/app.py  (بعد ضبط ADMIN_PANEL_USERNAME/ADMIN_PANEL_PASSWORD بملف .env)

import functools
import secrets
import sys
from pathlib import Path

# نضيف مجلد المشروع الرئيسي (اللي فيه db.py و config.py) لمسار البحث، لأن هذا الملف
# داخل مجلد فرعي (admin_panel/).
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import io
import os
import shutil

from flask import Flask, Response, abort, flash, redirect, render_template, request, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

import config
import countries
import db
import professions_repo

if not config.ADMIN_PANEL_PASSWORD:
    raise RuntimeError(
        "ADMIN_PANEL_PASSWORD غير موجود. أضف بملف .env:\n"
        "  ADMIN_PANEL_USERNAME=admin\n"
        "  ADMIN_PANEL_PASSWORD=كلمة_مرور_قوية_تختارها\n"
        "بدون كلمة مرور، أي شخص يعرف رابط اللوحة يقدر يدخلها — لذا نرفض التشغيل."
    )

app = Flask(__name__)
app.secret_key = config.ADMIN_PANEL_SECRET_KEY or secrets.token_hex(32)

PAGE_SIZE = 20

db.init_db()
db.seed_professions_from_json_if_empty(config.PROFESSIONS_JSON_PATH)
db.ensure_extra_professions()
db.seed_saudi_geo_if_empty(
    config.SAUDI_REGIONS_JSON_PATH, config.SAUDI_CITIES_JSON_PATH, config.SAUDI_DISTRICTS_JSON_PATH
)


# ─────────────────────────── تسجيل الدخول ───────────────────────────

# صلاحيتان:
#   super   — المدير العام (بيانات الدخول من ملف .env): كل شي.
#   country — مدير دولة (يُضاف من صفحة «المدراء»): يشوف ويدير فنيي دولته فقط، بدون
#             تنزيل إكسل، ولا إعدادات عامة، ولا تعديل المهن، ولا المدفوعات، ولا المدراء.

def login_required(view):
    @functools.wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login", next=request.full_path))
        return view(*args, **kwargs)
    return wrapped


def is_super() -> bool:
    return session.get("role", "super") == "super"


def super_required(view):
    """صفحات للمدير العام فقط — مدير الدولة يرجع للرئيسية مع تنبيه."""
    @functools.wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if not is_super():
            flash("هذي الصفحة متاحة للمدير العام فقط.", "error")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)
    return wrapped


def scope_country() -> str | None:
    """الدولة اللي تنحصر فيها البيانات: مدير الدولة دائمًا دولته (ما يقدر يغيّرها)،
    والمدير العام حسب فلتر ?country= (فاضي = كل الدول)."""
    if not is_super():
        return session.get("country")
    code = (request.args.get("country") or "").upper()
    return code if countries.get(code) else None


def _professional_or_404(professional_id: int):
    """يجيب الفني ويتأكد إن المستخدم الحالي له صلاحية عليه (مدير الدولة: دولته فقط)."""
    p = db.get_professional_by_id(professional_id)
    if not p:
        return None
    if not is_super() and (p.get("country") or "SA") != session.get("country"):
        abort(403)
    return p


@app.context_processor
def inject_role():
    return {
        "is_super": is_super() if session.get("logged_in") else False,
        "manager_country": session.get("country"),
        "country_label": countries.label,
        "all_countries": countries.COUNTRIES,
    }


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        next_url = request.args.get("next")
        if username == config.ADMIN_PANEL_USERNAME and password == config.ADMIN_PANEL_PASSWORD:
            session.clear()
            session["logged_in"] = True
            session["role"] = "super"
            flash("تم تسجيل الدخول بنجاح.", "success")
            return redirect(next_url or url_for("dashboard"))
        manager = db.get_admin_user_by_username(username)
        if manager and check_password_hash(manager["password_hash"], password):
            session.clear()
            session["logged_in"] = True
            session["role"] = "country"
            session["country"] = manager["country"]
            session["username"] = manager["username"]
            flash(f"أهلًا {manager['username']} — مدير {countries.label(manager['country'])}.", "success")
            return redirect(url_for("dashboard"))
        flash("اسم المستخدم أو كلمة المرور غير صحيحة.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    flash("تم تسجيل الخروج.", "success")
    return redirect(url_for("login"))


# ─────────────────────────── الرئيسية / الإحصائيات ───────────────────────────

@app.route("/")
@login_required
def dashboard():
    country = scope_country()
    stats = db.get_admin_stats(country=country)
    health = _server_health() if is_super() else None
    return render_template(
        "dashboard.html", stats=stats, health=health, selected_country=country,
        unresponsive=db.admin_unresponsive_professionals(3, country),
        notify_off=db.admin_notify_off_professionals(country),
        nudges_expired=db.admin_nudges_expired(365, country),
        active_page="dashboard",
    )


# ─────────────────────────── صحة السيرفر (للمدير العام) ───────────────────────────
# حدود الألوان: أخضر = مرتاح، أصفر = قربت (خطط للترقية)، أحمر = رقّ الآن.
PEAK_USERS_YELLOW, PEAK_USERS_RED = 30, 50       # مستخدمين بنفس الدقيقة
RAM_YELLOW, RAM_RED = 75, 90                      # ٪ ذاكرة مستخدمة
DISK_YELLOW, DISK_RED = 80, 90                    # ٪ مساحة مستخدمة
CPU_YELLOW, CPU_RED = 0.7, 1.0                    # متوسط الحمل لكل نواة
DB_YELLOW_MB, DB_RED_MB = 2048, 5120              # حجم قاعدة البيانات

_LEVEL_ORDER = {"green": 0, "yellow": 1, "red": 2}


def _level(value, yellow, red) -> str:
    if value is None:
        return "green"
    return "red" if value >= red else ("yellow" if value >= yellow else "green")


def _server_health() -> dict:
    items = []

    # الحمل على المعالج (متوسط آخر 5 دقائق ÷ عدد الأنوية)
    try:
        load5 = os.getloadavg()[1]
        cores = os.cpu_count() or 1
        per_core = load5 / cores
        items.append({"label": "المعالج (CPU)", "value": f"{per_core * 100:.0f}٪ (أنوية: {cores})",
                      "level": _level(per_core, CPU_YELLOW, CPU_RED)})
    except (OSError, AttributeError):
        pass

    # الذاكرة (من /proc/meminfo — بدون أي مكتبة إضافية)
    try:
        info = {}
        with open("/proc/meminfo", encoding="utf-8") as f:
            for line in f:
                k, v = line.split(":", 1)
                info[k] = int(v.strip().split()[0])  # kB
        total, avail = info["MemTotal"], info.get("MemAvailable", info.get("MemFree", 0))
        used_pct = (total - avail) * 100 / total
        items.append({"label": "الذاكرة (RAM)", "value": f"{used_pct:.0f}٪ من {total / 1024 / 1024:.1f} GB",
                      "level": _level(used_pct, RAM_YELLOW, RAM_RED)})
    except (OSError, KeyError, ValueError, ZeroDivisionError):
        pass

    # المساحة
    try:
        du = shutil.disk_usage(str(config.DB_PATH.parent))
        disk_pct = du.used * 100 / du.total
        items.append({"label": "المساحة (Disk)", "value": f"{disk_pct:.0f}٪ من {du.total / 1024 ** 3:.0f} GB",
                      "level": _level(disk_pct, DISK_YELLOW, DISK_RED)})
    except OSError:
        pass

    # حجم قاعدة البيانات
    try:
        size_mb = sum(
            os.path.getsize(f"{config.DB_PATH}{suffix}")
            for suffix in ("", "-wal") if os.path.exists(f"{config.DB_PATH}{suffix}")
        ) / 1024 / 1024
        items.append({"label": "حجم قاعدة البيانات", "value": f"{size_mb:.1f} MB",
                      "level": _level(size_mb, DB_YELLOW_MB, DB_RED_MB)})
    except OSError:
        pass

    peaks = db.activity_peaks()
    items.append({"label": "أعلى مستخدمين بنفس الدقيقة (آخر يوم)", "value": str(peaks["peak_day"]),
                  "level": _level(peaks["peak_day"], PEAK_USERS_YELLOW, PEAK_USERS_RED)})
    items.append({"label": "أعلى مستخدمين بنفس الدقيقة (آخر أسبوع)", "value": str(peaks["peak_week"]),
                  "level": _level(peaks["peak_week"], PEAK_USERS_YELLOW, PEAK_USERS_RED)})
    items.append({"label": "رسائل وضغطات آخر ساعة", "value": str(peaks["updates_last_hour"]), "level": "green"})

    overall = max((i["level"] for i in items), key=lambda lv: _LEVEL_ORDER[lv], default="green")
    return {"items": items, "overall": overall, "has_activity": peaks["has_data"]}


# ─────────────────────────── الفنيون ───────────────────────────

def _profession_options():
    """قائمة المجالات مع مهنها، لخيارات الفلترة بصفحة الفنيون."""
    domains = []
    for d in professions_repo.get_domains():
        domains.append({"id": d["id"], "name": d["name"], "professions": professions_repo.get_professions_by_domain(d["id"])})
    return domains


@app.route("/professionals")
@login_required
def professionals_list():
    status = request.args.get("status") or None
    city = request.args.get("city") or None
    profession_id = request.args.get("profession_id") or None
    subscribed_raw = request.args.get("subscribed") or None
    query = request.args.get("q") or None
    channel = request.args.get("channel") if request.args.get("channel") in ("whatsapp", "telegram") else None
    page = max(1, request.args.get("page", 1, type=int))

    subscribed = None
    if subscribed_raw == "yes":
        subscribed = True
    elif subscribed_raw == "no":
        subscribed = False

    country = scope_country()
    total_count = db.admin_count_professionals(
        status=status, city=city, profession_id=profession_id, subscribed=subscribed, query=query,
        country=country, channel=channel,
    )
    total_pages = max(1, (total_count + PAGE_SIZE - 1) // PAGE_SIZE)
    page = min(page, total_pages)

    professionals = db.admin_list_professionals(
        status=status, city=city, profession_id=profession_id, subscribed=subscribed, query=query,
        country=country, channel=channel, limit=PAGE_SIZE, offset=(page - 1) * PAGE_SIZE,
    )

    filters = {
        "status": status, "city": city, "profession_id": profession_id, "subscribed": subscribed_raw, "q": query,
        "country": country if is_super() else None, "channel": channel,
    }
    filters_qs = {k: v for k, v in filters.items() if v}

    return render_template(
        "professionals.html",
        professionals=professionals,
        total_count=total_count,
        page=page,
        total_pages=total_pages,
        filters=filters,
        filters_qs=filters_qs,
        cities=db.admin_list_cities(country),
        profession_options=_profession_options(),
        status_labels=db.STATUS_LABELS_AR,
        active_page="professionals",
    )


@app.route("/professionals/export.xlsx")
@super_required
def professionals_export():
    """تنزيل قائمة الفنيين (بنفس فلاتر صفحة الفنيون الحالية إن وُجدت) كملف إكسل."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font
    except ImportError:
        flash("ميزة التصدير تحتاج مكتبة openpyxl غير مثبتة على السيرفر — شغّل: pip install openpyxl", "error")
        return redirect(url_for("professionals_list"))

    status = request.args.get("status") or None
    city = request.args.get("city") or None
    profession_id = request.args.get("profession_id") or None
    subscribed_raw = request.args.get("subscribed") or None
    query = request.args.get("q") or None

    subscribed = None
    if subscribed_raw == "yes":
        subscribed = True
    elif subscribed_raw == "no":
        subscribed = False

    country = scope_country()
    total_count = db.admin_count_professionals(
        status=status, city=city, profession_id=profession_id, subscribed=subscribed, query=query,
        country=country,
    )
    professionals = db.admin_list_professionals(
        status=status, city=city, profession_id=profession_id, subscribed=subscribed, query=query,
        country=country, limit=max(total_count, 1), offset=0,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "الفنيون"
    ws.sheet_view.rightToLeft = True

    headers = [
        "الاسم", "رقم الواتساب", "رقم تواصل تيليجرام", "المهنة", "المجال",
        "الدولة", "المدينة", "الحي", "الحالة", "مشترك", "تاريخ انتهاء الاشتراك",
        "فرص مجانية مستخدمة", "مصدر التسجيل", "تاريخ التسجيل",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for p in professionals:
        ws.append([
            p["full_name"],
            p["whatsapp_number"],
            p["telegram_contact_number"] or "",
            db.profession_display(p),
            p["domain_name"],
            countries.name(p.get("country") or "SA"),
            p["city"],
            "المدينة كاملة" if p.get("covers_whole_city") else (p["neighborhood"] or ""),
            db.STATUS_LABELS_AR.get(p["status"], p["status"]),
            "نعم" if p["is_subscribed"] else "لا",
            (p["subscription_expires_at"] or "")[:10],
            p["free_contacts_used"],
            p.get("source") or "unknown",
            (p["created_at"] or "")[:16],
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # ملاحظة مهمة: اسم ملف بحروف عربية في Content-Disposition العادي (filename=...)
    # يكسر الترويسة (HTTP headers لازم تكون ASCII/latin-1) ويسبب خطأ 500 بدل ما
    # ينزل الملف — لازم نستخدم الصيغة المشفّرة filename*=UTF-8''... مع اسم احتياطي
    # إنجليزي بسيط بالصيغة العادية لأي متصفح قديم ما يفهم الصيغة الجديدة.
    from urllib.parse import quote

    encoded_name = quote("الفنيون.xlsx")
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f"attachment; filename=professionals.xlsx; filename*=UTF-8''{encoded_name}"
        },
    )


@app.route("/professionals/<int:professional_id>")
@login_required
def professional_detail(professional_id):
    p = _professional_or_404(professional_id)
    if not p:
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))
    return render_template(
        "professional_detail.html", p=p, status_labels=db.STATUS_LABELS_AR, active_page="professionals",
        notif=db.notification_summary(p["id"]), notif_stopped=_nudges_stopped(p),
        rating=__import__("ratings").summary([p["id"]]).get(p["id"]),
    )


def _nudges_stopped(p) -> dict:
    """حالة جدول إشعارات واتساب: الشهر الحالي، وهل توقفت (مرت سنة)."""
    import nudges
    st = nudges.wa_schedule_state(p)
    st["has_wa"] = bool(nudges.wa_number(p))
    return st


@app.route("/professionals/<int:professional_id>/notify-on", methods=["POST"])
@login_required
def professional_notify_on(professional_id):
    p = _professional_or_404(professional_id)
    if not p:
        return redirect(url_for("professionals_list"))
    ch = request.form.get("channel")
    if ch in ("telegram", "whatsapp"):
        db.set_notify_off(p["id"], ch, False)
        flash(f"تم تفعيل إشعارات {'تلغرام' if ch == 'telegram' else 'واتساب'} للفني ✅", "success")
    return redirect(request.referrer or url_for("professional_detail", professional_id=p["id"]))


@app.route("/professionals/<int:professional_id>/nudges-restart", methods=["POST"])
@login_required
def professional_nudges_restart(professional_id):
    p = _professional_or_404(professional_id)
    if not p:
        return redirect(url_for("professionals_list"))
    db.restart_wa_nudges(p["id"])
    flash("تم إعادة تشغيل إشعارات واتساب من البداية (3 إشعارات بالشهر الأول، ثم ملخص شهري لمدة سنة) ✅", "success")
    return redirect(request.referrer or url_for("professional_detail", professional_id=p["id"]))


@app.route("/professionals/<int:professional_id>/notify", methods=["POST"])
@login_required
def professional_notify(professional_id):
    p = _professional_or_404(professional_id)
    if not p:
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))
    import nudges
    ok = nudges.notify_manual(p["id"])
    via = "تلغرام" if (p.get("telegram_user_id") or 0) > 0 else "واتساب"
    flash(f"تم إرسال الإشعار عبر {via} ✅" if ok else f"ما قدرنا نرسل الإشعار عبر {via} (يمكن حاظر البوت، أو قالب واتساب ما انقبل بعد).",
          "success" if ok else "error")
    return redirect(request.referrer or url_for("professional_detail", professional_id=p["id"]))


@app.route("/professionals/<int:professional_id>/status", methods=["POST"])
@login_required
def professional_set_status(professional_id):
    if not _professional_or_404(professional_id):
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))
    new_status = request.form.get("status")
    if new_status not in (db.STATUS_PENDING, db.STATUS_ACTIVE, db.STATUS_REJECTED):
        flash("حالة غير معروفة.", "error")
    else:
        db.set_status(professional_id, new_status)
        flash("تم تحديث حالة الفني.", "success")
    next_url = request.form.get("next")
    return redirect(next_url or url_for("professional_detail", professional_id=professional_id))


@app.route("/professionals/<int:professional_id>/edit", methods=["POST"])
@login_required
def professional_edit(professional_id):
    p = _professional_or_404(professional_id)
    if not p:
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))

    # نرتّب الأرقام لصيغة دولية حسب دولة الفني (مثل 01012345678 ← +201012345678) —
    # بدونه رابط واتساب للعميل يطلع على رقم غلط. رقم ما ينفهم → نرفضه وما نحفظه.
    country = p.get("country") or countries.DEFAULT_COUNTRY
    wa_raw = (request.form.get("whatsapp_number") or "").strip()
    tg_raw = (request.form.get("telegram_contact_number") or "").strip()
    wa = countries.normalize_phone(wa_raw, country) if wa_raw else None
    if wa_raw and not wa:
        c = countries.get(country)
        flash(f"رقم الواتساب غير صحيح لـ{countries.name(country)} — مثال صحيح: {c['example']}", "error")
        return redirect(url_for("professional_detail", professional_id=professional_id))
    tg = (countries.normalize_phone(tg_raw, country) or tg_raw) if tg_raw else ""

    db.admin_update_professional(
        professional_id,
        full_name=request.form.get("full_name") or None,
        city=request.form.get("city") or None,
        neighborhood=request.form.get("neighborhood") or "",
        whatsapp_number=wa,
        telegram_contact_number=tg,
        has_whatsapp=request.form.get("has_whatsapp") == "1",
    )
    flash("تم حفظ بيانات الفني.", "success")
    return redirect(url_for("professional_detail", professional_id=professional_id))


@app.route("/professionals/<int:professional_id>/subscription", methods=["POST"])
@login_required
def professional_subscription(professional_id):
    if not _professional_or_404(professional_id):
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))
    action = request.form.get("action")
    expires_at = request.form.get("expires_at") or None
    if action == "activate":
        db.admin_set_subscription(professional_id, True, expires_at)
        flash("تم تفعيل الاشتراك يدويًا.", "success")
    elif action == "deactivate":
        db.admin_set_subscription(professional_id, False, None)
        flash("تم إلغاء الاشتراك.", "success")
    return redirect(url_for("professional_detail", professional_id=professional_id))


@app.route("/professionals/<int:professional_id>/reset-free-contacts", methods=["POST"])
@login_required
def professional_reset_free(professional_id):
    if not _professional_or_404(professional_id):
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))
    db.admin_reset_free_contacts(professional_id)
    flash("تم تصفير عداد الفرص المجانية.", "success")
    return redirect(url_for("professional_detail", professional_id=professional_id))


@app.route("/professionals/<int:professional_id>/delete", methods=["POST"])
@login_required
def professional_delete(professional_id):
    if not _professional_or_404(professional_id):
        flash("الفني غير موجود.", "error")
        return redirect(url_for("professionals_list"))
    db.admin_delete_professional(professional_id)
    flash("تم حذف الفني نهائيًا.", "success")
    return redirect(url_for("professionals_list"))


# ─────────────────────────── مؤشرات التغطية (جاهزية الإطلاق) ───────────────────────────

# أي مدينة فيها "عدد كافٍ" من الفنيين لمهنة معينة تُعتبر مغطّاة بهالمهنة — حد أدنى
# بسيط قابل للتعديل لاحقًا حسب الحاجة الفعلية.
COVERAGE_READY_THRESHOLD = 3


@app.route("/coverage")
@login_required
def coverage_page():
    # مدن الدول المختلفة ممكن تتشابه أسماؤها، فالتغطية دائمًا لدولة وحدة (الافتراضي السعودية)
    country = scope_country() or countries.DEFAULT_COUNTRY
    summary = db.admin_coverage_summary(country)
    matrix_rows = db.admin_coverage_matrix(country)

    # counts[profession_id][city] = عدد الفنيين النشطين
    counts: dict[str, dict[str, int]] = {}
    for r in matrix_rows:
        counts.setdefault(r["profession_id"], {})[r["city"]] = r["cnt"]

    domains = []
    total_professions = 0
    for d in professions_repo.get_domains():
        profs = professions_repo.get_professions_by_domain(d["id"])
        domains.append({"id": d["id"], "name": d["name"], "professions": profs})
        total_professions += len(profs)

    # فلتر المدينة: القائمة فيها بس المدن اللي فيها فنيين نشطين فعلًا بهالدولة
    all_cities = [row["city"] for row in summary]
    selected_city = request.args.get("city") or None
    if selected_city not in all_cities:
        selected_city = None
    if selected_city:
        summary = [row for row in summary if row["city"] == selected_city]

    # عدد المهن "الجاهزة" (وصلت الحد الأدنى) لكل مدينة — لحساب نسبة الجاهزية
    cities = [row["city"] for row in summary]
    ready_counts = {city: 0 for city in cities}
    for prof_counts in counts.values():
        for city, cnt in prof_counts.items():
            if city in ready_counts and cnt >= COVERAGE_READY_THRESHOLD:
                ready_counts[city] += 1

    return render_template(
        "coverage.html",
        summary=summary,
        cities=cities,
        counts=counts,
        domains=domains,
        total_professions=total_professions,
        ready_counts=ready_counts,
        ready_threshold=COVERAGE_READY_THRESHOLD,
        selected_country=country,
        all_cities=all_cities,
        selected_city=selected_city,
        active_page="coverage",
    )


# ─────────────────────────── الاشتراكات والمدفوعات ───────────────────────────

@app.route("/payments")
@super_required
def payments_list():
    status = request.args.get("status") or None
    method = request.args.get("method") or None
    page = max(1, request.args.get("page", 1, type=int))

    total_count = db.admin_count_payments(status=status, method=method)
    total_pages = max(1, (total_count + PAGE_SIZE - 1) // PAGE_SIZE)
    page = min(page, total_pages)

    payments = db.admin_list_payments(status=status, method=method, limit=PAGE_SIZE, offset=(page - 1) * PAGE_SIZE)

    filters = {"status": status, "method": method}
    filters_qs = {k: v for k, v in filters.items() if v}

    return render_template(
        "payments.html",
        payments=payments,
        total_count=total_count,
        page=page,
        total_pages=total_pages,
        filters=filters,
        filters_qs=filters_qs,
        active_page="payments",
    )


# ─────────────────────────── الإعدادات العامة ───────────────────────────

SETTINGS_DEFAULTS = {
    "subscription_price_sar": str(config.SUBSCRIPTION_PRICE_SAR),
    "subscription_price_stars": str(config.SUBSCRIPTION_PRICE_STARS),
    "subscription_days": str(config.SUBSCRIPTION_DAYS),
    "free_contacts_limit": str(db.FREE_CONTACTS_LIMIT),
}


@app.route("/settings", methods=["GET", "POST"])
@super_required
def settings_page():
    if request.method == "POST":
        for key in SETTINGS_DEFAULTS:
            value = request.form.get(key)
            if value:
                db.set_setting(key, value)
        enabled = [code for code in countries.ALL_CODES if request.form.get(f"country_{code}") == "1"]
        if countries.DEFAULT_COUNTRY not in enabled:
            enabled.insert(0, countries.DEFAULT_COUNTRY)
        db.set_setting("enabled_countries", ",".join(enabled))
        # لكل دولة: أي بوت تظهر فيه، وأي طرق دفع متاحة لها
        for key, field in (("wa_countries", "wa"), ("tg_countries", "tg"),
                           ("tap_countries", "tap"), ("stars_countries", "stars")):
            db.set_setting(key, ",".join(c for c in countries.ALL_CODES if request.form.get(f"{field}_{c}") == "1"))
        warn = []
        for c in enabled:
            label = countries.name(c, "ar")
            if request.form.get(f"wa_{c}") != "1" and request.form.get(f"tg_{c}") != "1":
                warn.append(f"{label}: متاحة لكن ما اخترت لها واتساب ولا تلغرام")
            if request.form.get(f"tap_{c}") != "1" and request.form.get(f"stars_{c}") != "1":
                warn.append(f"{label}: ما فيها أي طريقة دفع — فنيينها ما يقدرون يشتركون")
        flash("تم حفظ الإعدادات — تنعكس فورًا على البوت.", "success")
        for w in warn:
            flash("⚠️ " + w, "error")
        return redirect(url_for("settings_page"))

    current = {key: db.get_setting(key, default) for key, default in SETTINGS_DEFAULTS.items()}
    sets = {k: countries._codes(k, d) for k, d in (
        ("enabled_countries", ",".join(countries.ALL_CODES)), ("wa_countries", ",".join(countries.ALL_CODES)),
        ("tg_countries", ",".join(countries.ALL_CODES)), ("tap_countries", countries.PAYMENT_DEFAULTS["tap_countries"]),
        ("stars_countries", countries.PAYMENT_DEFAULTS["stars_countries"]))}
    return render_template(
        "settings.html", settings=current, enabled_countries=countries.enabled_codes(), country_sets=sets,
        active_page="settings",
    )


# ─────────────────────────── مدراء الدول ───────────────────────────

@app.route("/admins", methods=["GET", "POST"])
@super_required
def admins_page():
    if request.method == "POST":
        action = request.form.get("action")
        if action == "add":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            country = (request.form.get("country") or "").upper()
            if not username or len(password) < 6 or not countries.get(country):
                flash("اكتب اسم مستخدم، وكلمة مرور 6 أحرف على الأقل، واختر الدولة.", "error")
            elif username == config.ADMIN_PANEL_USERNAME:
                flash("هذا الاسم محجوز للمدير العام — اختر اسم ثاني.", "error")
            elif db.create_admin_user(username, generate_password_hash(password), country):
                flash(f"تمت إضافة المدير «{username}» لـ {countries.label(country)}.", "success")
            else:
                flash("اسم المستخدم مستخدم مسبقًا.", "error")
        elif action == "password":
            password = request.form.get("password", "")
            if len(password) < 6:
                flash("كلمة المرور لازم تكون 6 أحرف على الأقل.", "error")
            else:
                db.update_admin_user_password(int(request.form["user_id"]), generate_password_hash(password))
                flash("تم تغيير كلمة المرور.", "success")
        elif action == "delete":
            db.delete_admin_user(int(request.form["user_id"]))
            flash("تم حذف المدير.", "success")
        return redirect(url_for("admins_page"))

    return render_template("admins.html", admins=db.list_admin_users(), active_page="admins")


# ─────────────────────────── المسوّقون (تسجيل فنيين على أرقام أخرى) ───────────────────────────

@app.route("/registrars", methods=["GET", "POST"])
@super_required
def registrars_page():
    if request.method == "POST":
        if request.form.get("action") == "otp":
            db.set_setting("wa_otp_enabled", "1" if request.form.get("on") == "1" else "0")
            flash("تم الحفظ.", "success")
        elif request.form.get("action") == "remove":
            db.remove_wa_registrar(request.form.get("phone", ""))
            flash("تم حذف الرقم.", "success")
        elif db.add_wa_registrar(request.form.get("phone", ""), request.form.get("name", "")):
            flash("تمت إضافة الرقم — يقدر الآن يسجّل فنيين بأرقامهم من بوت واتساب بدون رمز تحقق.", "success")
        else:
            flash("اكتب الرقم كاملًا مع رمز الدولة، مثل: 966501234567", "error")
        return redirect(url_for("registrars_page"))
    return render_template("registrars.html", rows=db.list_wa_registrars(), active_page="registrars",
                           otp_on=db.get_setting("wa_otp_enabled", "0") == "1")


# ─────────────────────────── الفنيون المحذوفون (أرشيف للتواصل) ───────────────────────────

@app.route("/deleted", methods=["GET", "POST"])
@super_required
def deleted_page():
    if request.method == "POST":
        aid = int(request.form.get("id") or 0)
        action = request.form.get("action")
        if action == "days":
            try:
                n = max(0, min(db.ARCHIVE_DAYS_MAX, int(request.form.get("days") or 0)))
            except ValueError:
                n = 1
            db.set_setting("archive_days", str(n))
            purged = db.purge_expired_archive()
            flash(f"تم الحفظ: {('يُحذف فورًا' if n == 0 else 'مدة الاحتفاظ ' + db.days_text(n))}. "
                  f"وتغيّرت رسالة الحذف وصفحة الخصوصية تلقائيًا." + (f" حُذف {purged} سجلًا انتهت مدته." if purged else ""),
                  "success")
        elif action == "save":
            db.set_deletion_reason(aid, request.form.get("reason", ""), "admin")
            db.set_deletion_note(aid, request.form.get("note", ""))
            flash("تم الحفظ.", "success")
        elif action == "erase":
            db.remove_deleted_archive(aid)
            flash("تم حذف السجل من الأرشيف نهائيًا.", "success")
        return redirect(url_for("deleted_page"))
    rows = db.list_deleted_professionals()
    with db.get_conn() as conn:
        back = {db._phone_digits(r[0] or "") for r in conn.execute("SELECT whatsapp_number FROM professionals").fetchall()}
    for r in rows:
        r["returned"] = r.get("phone") in back
    return render_template("deleted.html", rows=rows, active_page="deleted", days=db.archive_days())


def _riyadh(iso: str | None) -> str:
    """تاريخ ووقت بتوقيت الرياض (المخزّن UTC)."""
    from datetime import datetime, timedelta
    try:
        return (datetime.fromisoformat(iso) + timedelta(hours=3)).strftime("%Y-%m-%d %H:%M")
    except (TypeError, ValueError):
        return (iso or "")[:16]


@app.route("/deleted/export.xlsx")
@super_required
def deleted_export():
    """قائمة المحذوفين كملف إكسل — لخدمة العملاء للتواصل معهم عن سبب الحذف."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from urllib.parse import quote

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المحذوفون"
    ws.sheet_view.rightToLeft = True
    ws.append(["تاريخ الحذف", "الاسم", "رقم الواتساب", "رابط المحادثة", "المهنة", "المدينة", "سجّل من",
               "من حذفه", "التقييم", "عدد التقييمات", "سبب الحذف", "كتب السبب", "ملاحظات الإدارة",
               "نتيجة التواصل (لخدمة العملاء)"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F766E")
        cell.alignment = Alignment(horizontal="center")
    for r in db.list_deleted_professionals(limit=100000):
        phone = r.get("phone") or ""
        ws.append([
            _riyadh(r.get("deleted_at")),
            r.get("full_name") or "",
            ("+" + phone) if phone else "",
            f"https://wa.me/{phone}" if phone else "",
            r.get("profession_name") or "",
            r.get("city") or "",
            "واتساب" if r.get("channel") == "wa" else "تلغرام",
            "المشرف" if r.get("deleted_by") == "admin" else "الفني نفسه",
            r.get("rating_avg") if r.get("rating_count") else "",
            r.get("rating_count") or 0,
            r.get("reason") or "",
            {"pro": "الفني", "admin": "المشرف"}.get(r.get("reason_by") or "", ""),
            r.get("admin_note") or "",
            "",
        ])
        link = ws.cell(row=ws.max_row, column=4)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
    widths = [17, 22, 16, 30, 22, 14, 10, 12, 9, 12, 30, 10, 30, 32]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    name = quote("المحذوفون.xlsx")
    return Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename=deleted.xlsx; filename*=UTF-8''{name}"})


@app.route("/professionals/<int:professional_id>/reset-rating", methods=["POST"])
@super_required
def professional_reset_rating(professional_id):
    n = __import__("ratings").reset_for_professional(professional_id)
    flash(f"تم تصفير التقييم ({n} تقييم أُلغي). يبدأ الفني بتقييم جديد.", "success")
    return redirect(url_for("professional_detail", professional_id=professional_id))


# ─────────────────────────── المهن الأكثر طلبًا + قائمة الـ9 بواتساب ───────────────────────────

@app.route("/demand", methods=["GET", "POST"])
@super_required
def demand_page():
    from whatsapp_bot import top
    all_profs = top.all_professions()
    valid = {p["id"] for p in all_profs}
    if request.method == "POST":
        if request.form.get("auto") == "1":
            demand = db.get_profession_demand(int(request.form.get("days") or 30))
            ranked = sorted((p for p in all_profs if demand.get(p["id"], {}).get("recent", 0) > 0),
                            key=lambda p: demand[p["id"]]["recent"], reverse=True)
            chosen = [p["id"] for p in ranked[:9]]
            # لو المهن المطلوبة أقل من 9، نكمّل من القائمة الحالية (بدل مهن ما أحد بحث عنها)
            for p in top.top_professions(all_profs):
                if len(chosen) < 9 and p["id"] not in chosen:
                    chosen.append(p["id"])
        else:
            chosen = []
            for i in range(9):
                pid = request.form.get(f"slot{i}", "")
                if pid in valid and pid not in chosen:
                    chosen.append(pid)
        if len(chosen) != 9:
            flash("لازم تختار 9 مهن مختلفة.", "error")
        else:
            db.set_setting(top.SETTING_KEY, ",".join(chosen))
            flash("تم حفظ قائمة الـ9 — تظهر فورًا بواتساب.", "success")
        return redirect(url_for("demand_page", days=request.form.get("days") or 30))

    days = request.args.get("days", 30, type=int)
    if days not in (7, 30, 90, 365):
        days = 30
    demand = db.get_profession_demand(days)
    top_ids = [p["id"] for p in top.top_professions(all_profs)]
    rows = []
    for p in all_profs:
        d = demand.get(p["id"], {})
        rows.append({"id": p["id"], "name": p["name"], "recent": d.get("recent", 0), "total": d.get("total", 0),
                     "empty": d.get("empty", 0), "techs": d.get("techs", 0),
                     "top_pos": top_ids.index(p["id"]) + 1 if p["id"] in top_ids else None})
    rows.sort(key=lambda r: (r["recent"], r["total"]), reverse=True)
    for i, r in enumerate(rows, 1):
        r["rank"] = i
    # تنبيه: مهن داخل الـ9 مو ضمن أعلى 9 طلبًا (لما يكون فيه بيانات كافية)
    has_data = any(r["recent"] for r in rows)
    enough = sum(r["recent"] for r in rows) >= 30   # الاقتراح يحتاج بحث كافي عشان ما يضلّل
    top9_by_demand = {r["id"] for r in rows[:9]}
    suggest_out = [r for r in rows if enough and r["top_pos"] and r["id"] not in top9_by_demand]
    suggest_in = [r for r in rows[:9] if enough and not r["top_pos"] and r["recent"] > 0]
    options = sorted(all_profs, key=lambda p: p["name"])
    terms = db.get_search_terms(days)
    return render_template("demand.html", rows=rows, top_ids=top_ids, options=options, days=days,
                           terms_missing=[t for t in terms if not t["exact"]],
                           terms_found=[t for t in terms if t["exact"]],
                           suggest_out=suggest_out, suggest_in=suggest_in, has_data=has_data,
                           active_page="demand")


@app.route("/demand/ignore", methods=["POST"])
@super_required
def demand_ignore_term():
    key = (request.form.get("term_norm") or "").strip()
    if key:
        db.ignore_search_term(key)
        flash("تم تجاهل الكلمة — ما تطلع لك مرة ثانية.", "success")
    return redirect(url_for("demand_page", days=request.form.get("days") or 30) + "#terms")


# ─────────────────────────── إدارة المهن والمجالات ───────────────────────────

@app.route("/professions")
@super_required
def professions_page():
    return render_template("professions.html", domains=_profession_options(), active_page="professions")


@app.route("/professions/domains/add", methods=["POST"])
@super_required
def domain_add():
    name = request.form.get("name", "").strip()
    name_en = request.form.get("name_en", "").strip()
    name_ur = request.form.get("name_ur", "").strip()
    if name and (not name_en or not name_ur):
        flash("أضف ترجمة اسم المجال بالإنجليزي والأردو — البوت يعمل بهذه اللغات أيضًا.", "error")
        return redirect(url_for("professions_page"))
    if name:
        did = db.create_domain(name)
        db.set_domain_translations(did, name_en, name_ur)
        flash("تمت إضافة المجال.", "success")
    return redirect(url_for("professions_page"))


@app.route("/professions/domains/<domain_id>/rename", methods=["POST"])
@super_required
def domain_rename(domain_id):
    name = request.form.get("name", "").strip()
    if name:
        db.rename_domain(domain_id, name)
        if request.form.get("name_en") is not None:
            db.set_domain_translations(domain_id, request.form.get("name_en"), request.form.get("name_ur"))
        flash("تم تعديل اسم المجال.", "success")
    return redirect(url_for("professions_page"))


@app.route("/professions/domains/<domain_id>/delete", methods=["POST"])
@super_required
def domain_delete(domain_id):
    if db.delete_domain(domain_id):
        flash("تم حذف المجال.", "success")
    else:
        flash("لا يمكن حذف مجال يحتوي على مهن — احذف/انقل مهنه أولًا.", "error")
    return redirect(url_for("professions_page"))


def _parse_services(raw: str) -> list:
    return [s.strip() for s in raw.split("،" if "،" in raw else ",") if s.strip()]


@app.route("/professions/domains/<domain_id>/professions/add", methods=["POST"])
@super_required
def profession_add(domain_id):
    name = request.form.get("name", "").strip()
    isco_code = request.form.get("isco_code", "").strip() or None
    services = _parse_services(request.form.get("services", ""))
    name_en = request.form.get("name_en", "").strip()
    name_ur = request.form.get("name_ur", "").strip()
    if name and (not name_en or not name_ur):
        flash("أضف ترجمة اسم المهنة بالإنجليزي والأردو — البوت يعمل بهذه اللغات أيضًا.", "error")
        return redirect(url_for("professions_page"))
    if name:
        pid = db.create_profession(domain_id, name, isco_code, services)
        db.set_profession_translations(pid, name_en, name_ur)
        flash("تمت إضافة المهنة.", "success")
    return redirect(url_for("professions_page"))


@app.route("/professions/<profession_id>/edit", methods=["GET", "POST"])
@super_required
def profession_edit_page(profession_id):
    prof = db.get_profession_by_id(profession_id)
    if not prof:
        flash("المهنة غير موجودة.", "error")
        return redirect(url_for("professions_page"))

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        isco_code = request.form.get("isco_code", "").strip() or None
        services = _parse_services(request.form.get("services", ""))
        allow_city_wide = request.form.get("allow_city_wide") == "1"
        name_en = request.form.get("name_en", "").strip()
        name_ur = request.form.get("name_ur", "").strip()
        if name and (not name_en or not name_ur):
            flash("أضف ترجمة اسم المهنة بالإنجليزي والأردو — البوت يعمل بهذه اللغات أيضًا.", "error")
            return redirect(url_for("profession_edit_page", profession_id=profession_id))
        if name:
            db.update_profession(profession_id, name, isco_code, services, allow_city_wide=allow_city_wide)
            db.set_profession_translations(profession_id, name_en, name_ur)
            flash("تم حفظ تعديلات المهنة.", "success")
            return redirect(url_for("professions_page"))
        flash("اسم المهنة مطلوب.", "error")

    import json as _json
    prof_view = {
        "id": prof["id"],
        "name": prof["name"],
        "isco_code": prof.get("isco_code"),
        "services": _json.loads(prof["services_json"]) if prof.get("services_json") else [],
        "allow_city_wide": bool(prof.get("allow_city_wide")),
        "name_en": prof.get("name_en") or "",
        "name_ur": prof.get("name_ur") or "",
    }
    return render_template("profession_edit.html", profession=prof_view, active_page="professions")


@app.route("/professions/<profession_id>/delete", methods=["POST"])
@super_required
def profession_delete(profession_id):
    if db.delete_profession(profession_id):
        flash("تم حذف المهنة.", "success")
    else:
        flash("لا يمكن حذف مهنة مرتبط بها فنيون مسجّلون حاليًا.", "error")
    return redirect(url_for("professions_page"))


# ─────────────────────────── تحويل أزرار التواصل (عام، بدون تسجيل دخول) ───────────────────────────
# زر «تواصل عبر واتساب/تلغرام» ببطاقة الفني يفتح هذا الرابط: نسجّل التواصل (لحساب فرص
# الفني المجانية) ثم نحوّل العميل فورًا لواتساب/تلغرام — ضغطة وحدة، والرقم ما يظهر
# بالبطاقة. الرمز موقّع (contact_links) فمحد يقدر يصنع رابط لفني ثاني أو عميل ثاني.

import contact_links  # noqa: E402
import i18n  # noqa: E402


def _bot_link_page(lang: str, status: int = 410):
    rtl = lang in ("ar", "ur")
    html = f"""<!doctype html><html lang="{lang}" dir="{'rtl' if rtl else 'ltr'}"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>فني</title>
<style>body{{font-family:system-ui,Tahoma,sans-serif;background:#0B1626;color:#fff;display:flex;min-height:100vh;
align-items:center;justify-content:center;margin:0;padding:24px;text-align:center}}
a{{display:inline-block;margin-top:22px;background:#2EC4B6;color:#0B1626;padding:14px 28px;border-radius:999px;
font-weight:700;text-decoration:none}}</style></head><body><div><div style="font-size:48px">🛠️</div>
<p style="font-size:19px;line-height:1.7;max-width:420px">{i18n.t("link_expired_page", lang)}</p>
<a href="https://t.me/FanniServiceBot">{i18n.t("link_back_to_bot", lang)}</a></div></body></html>"""
    return Response(html, status=status, mimetype="text/html")


def _maybe_donation_prompt(customer_id: int, lang: str):
    """نفس رسالة الدعم الاختيارية اللي كان البوت يرسلها كل 10 تواصلات — ترسل بالخلفية
    حتى ما يتأخر تحويل العميل لواتساب. أي فشل هنا ما يأثر على التحويل."""
    try:
        from handlers import donation
        count = db.count_contact_clicks_by_customer(customer_id)
        if count == 0 or count % donation.DONATE_EVERY_N_CONTACTS != 0:
            return
        import asyncio
        import threading

        from telegram import Bot

        async def _send():
            async with Bot(config.BOT_TOKEN) as bot:
                await bot.send_message(
                    chat_id=customer_id,
                    text=i18n.t("donate_prompt", lang, count=count),
                    reply_markup=donation._amount_keyboard(lang),
                )

        threading.Thread(target=lambda: asyncio.run(_send()), daemon=True).start()
    except Exception:
        app.logger.exception("donation prompt failed")


@app.route("/c/<token>")
def contact_redirect(token):
    data = contact_links.parse_token(token)
    if not data:
        return _bot_link_page("ar")
    cid = data["customer_id"]
    lang = db.get_user_language(cid) or "ar"
    p = db.get_professional_by_id(data["professional_id"])
    if not p or p.get("status") == db.STATUS_REJECTED:
        return _bot_link_page(lang)

    if data["channel"] == "w" and p.get("has_whatsapp", 1):
        prefill = i18n.t("srch_wa_prefill_text", lang, profession=contact_links.profession_for_slot(p, data["prof_slot"]))
        target = contact_links.wa_link(p["whatsapp_number"], prefill, p.get("country"))
    elif data["channel"] == "t" and p.get("telegram_contact_number"):
        target = contact_links.tg_link(p["telegram_contact_number"], p.get("country"))
    else:
        return _bot_link_page(lang)

    if db.register_contact(p["id"], cid):
        _maybe_donation_prompt(cid, lang)
    resp = redirect(target, code=302)
    resp.headers["Cache-Control"] = "no-store"
    return resp


if __name__ == "__main__":
    import os as _os

    # الافتراضي 127.0.0.1 (محلي فقط، للتطوير) — على السيرفر نضبط ADMIN_PANEL_HOST=0.0.0.0
    # عبر متغير بيئة (systemd Environment=) حتى يكون الرابط قابل للفتح من المتصفح مباشرة.
    host = _os.environ.get("ADMIN_PANEL_HOST", "127.0.0.1")
    port = int(_os.environ.get("ADMIN_PANEL_PORT", "5050"))
    app.run(host=host, port=port, debug=False)
