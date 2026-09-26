# -*- coding: utf-8 -*-
"""
يحوّل ملف تصنيف المهنيين (xlsx) إلى professions.json
البنية: قائمة مجالات، كل مجال فيه قائمة مهن، كل مهنة فيها (اسم، رمز ISCO، خدمات فرعية).
هذا الملف هو مصدر البيانات الوحيد المشترك بين /register و /search.
"""
import json
import openpyxl
from pathlib import Path

SRC = Path(__file__).parent / "professions_source.xlsx"
OUT = Path(__file__).parent / "professions.json"

NO_SERVICES_MARKERS = {None, "", "لا توجد خدمات فرعية مذكورة", "NaN"}


def parse_services(raw):
    if raw in NO_SERVICES_MARKERS:
        return []
    # الخدمات مفصولة بفاصلة عربية "، " في الملف الأصلي
    parts = [p.strip() for p in str(raw).replace("،", ",").split(",")]
    return [p for p in parts if p]


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["المهنيون والفنيون (ISCO-08)"]

    headers = [ws.cell(row=1, column=c).value for c in range(1, 8)]
    assert headers[0] == "المجال", f"ترتيب الأعمدة تغيّر: {headers}"

    domains = {}  # اسم المجال -> قائمة مهن
    domain_order = []  # نحافظ على ترتيب ظهور المجالات بالملف

    for row in range(2, ws.max_row + 1):
        domain = ws.cell(row=row, column=1).value
        subcat = ws.cell(row=row, column=2).value
        isco = ws.cell(row=row, column=3).value
        official_en = ws.cell(row=row, column=4).value
        local_name = ws.cell(row=row, column=5).value
        status = ws.cell(row=row, column=6).value
        services_raw = ws.cell(row=row, column=7).value

        if not domain or not local_name:
            continue  # صف فاضي أو ناقص

        if domain not in domains:
            domains[domain] = []
            domain_order.append(domain)

        # معرّف ثابت للمهنة (slug) يُستخدم بأزرار الـ callback_data
        slug = f"p{row-1}"

        domains[domain].append({
            "id": slug,
            "name": local_name,
            "subcategory": subcat,
            "isco_code": isco,
            "official_name_en": official_en,
            "status": status,
            "services": parse_services(services_raw),
        })

    result = {
        "domains": [
            {
                "id": f"d{i+1}",
                "name": domain_name,
                "professions": domains[domain_name],
            }
            for i, domain_name in enumerate(domain_order)
        ]
    }

    total_professions = sum(len(d["professions"]) for d in result["domains"])

    OUT.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"تم إنشاء {OUT} — {len(result['domains'])} مجال، {total_professions} مهنة")


if __name__ == "__main__":
    main()
